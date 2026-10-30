
# import sqlite3
# import tempfile
# from pathlib import Path

# from src.database.connection import get_connection
# from src.database.review_manager import review_record
# from src.review.export_verified_data import (
#     export_verified_data,
#     get_verified_records
# )


# with tempfile.TemporaryDirectory() as temp_dir:
#     temp_dir = Path(temp_dir)

#     source = get_connection()
#     connection = sqlite3.connect(temp_dir / "test.db")

#     try:
#         source.backup(connection)
#     finally:
#         source.close()

#     connection.row_factory = sqlite3.Row
#     connection.execute("PRAGMA foreign_keys = ON")

#     try:
#         records = connection.execute(
#             """
#             SELECT id, count
#             FROM extraction_records
#             WHERE review_status = 'pending'
#             ORDER BY id
#             LIMIT 3
#             """
#         ).fetchall()

#         if len(records) < 3:
#             raise ValueError("Need 3 pending records")

#         with connection:
#             review_record(
#                 connection,
#                 records[0]["id"],
#                 "approved",
#                 "Test Reviewer"
#             )

#             review_record(
#                 connection,
#                 records[1]["id"],
#                 "corrected",
#                 "Test Reviewer",
#                 corrections={
#                     "count": records[1]["count"] + 1
#                 }
#             )

#             review_record(
#                 connection,
#                 records[2]["id"],
#                 "rejected",
#                 "Test Reviewer"
#             )

#         verified = get_verified_records(connection)

#         exported_ids = {
#             record["record_id"] for record in verified
#         }

#         assert records[0]["id"] in exported_ids
#         assert records[1]["id"] in exported_ids
#         assert records[2]["id"] not in exported_ids

#         corrected = next(
#             record for record in verified
#             if record["record_id"] == records[1]["id"]
#         )

#         assert corrected["count"] == records[1]["count"] + 1

#         total = export_verified_data(
#             temp_dir / "output",
#             connection
#         )

#         assert total == 2

#         print("\nAll verified export tests passed!")

#     finally:
#         connection.close()


import tempfile
from pathlib import Path
from openpyxl import load_workbook

from src.review.export_verified_data import (
    export_verified_data,
    read_verified_records
)


SOURCE = Path("data/review/cieh_review.xlsx")

with tempfile.TemporaryDirectory() as directory:
    temp_dir = Path(directory)
    test_file = temp_dir / "test_review.xlsx"

    workbook = load_workbook(SOURCE)
    sheet = workbook["Review Records"]

    headers = {
        cell.value: cell.column
        for cell in sheet[1]
    }

    if sheet.max_row < 4:
        raise ValueError("Need at least three records")

    # Clear any existing review decisions in the test copy.
    for row in range(2, sheet.max_row + 1):
        for name in [
            "Review Decision", "Corrected Year",
            "Corrected Category", "Corrected Metric Type",
            "Corrected Count", "Reviewer", "Reviewer Notes"
        ]:
            sheet.cell(row, headers[name]).value = None

    # Approve first record
    sheet.cell(2, headers["Review Decision"]).value = "approved"
    sheet.cell(2, headers["Reviewer"]).value = "Test Reviewer"

    # Correct second record
    original_count = sheet.cell(3, headers["Count"]).value

    sheet.cell(3, headers["Review Decision"]).value = "corrected"
    sheet.cell(3, headers["Corrected Count"]).value = original_count + 1
    sheet.cell(3, headers["Reviewer"]).value = "Test Reviewer"

    # Reject third record
    sheet.cell(4, headers["Review Decision"]).value = "rejected"
    sheet.cell(4, headers["Reviewer"]).value = "Test Reviewer"

    workbook.save(test_file)
    workbook.close()

    records = read_verified_records(test_file)

    assert len(records) == 2
    assert records[0]["Review Decision"] == "approved"
    assert records[1]["Review Decision"] == "corrected"
    assert records[1]["Count"] == original_count + 1

    total = export_verified_data(
        test_file,
        temp_dir / "output"
    )

    assert total == 2
    assert (temp_dir / "output" / "cieh_verified.csv").exists()
    assert (temp_dir / "output" / "cieh_verified.xlsx").exists()

    print("\nAll database-free export tests passed!")
