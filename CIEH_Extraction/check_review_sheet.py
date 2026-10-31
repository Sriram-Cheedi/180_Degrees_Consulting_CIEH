from pathlib import Path
from collections import Counter
from openpyxl import load_workbook

path = Path("data/review/cieh_review.xlsx")

workbook = load_workbook(path, read_only=True, data_only=True)

try:
    sheet = workbook["Review Records"]

    headers = [cell.value for cell in next(sheet.iter_rows(max_row=1))]

    decision_index = headers.index("Review Decision")
    decisions = Counter()

    for row in sheet.iter_rows(min_row=2, values_only=True):
        decision = row[decision_index]

        if decision is None or str(decision).strip() == "":
            decisions["pending"] += 1
        else:
            decisions[str(decision).strip().lower()] += 1

    print("=== EXCEL REVIEW SUMMARY ===")

    for decision, count in sorted(decisions.items()):
        print(f"{decision}: {count}")

finally:
    workbook.close()
