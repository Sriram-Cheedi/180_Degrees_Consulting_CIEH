
from pathlib import Path
from openpyxl import load_workbook

from src.ai.extraction_schema import ExtractionRecord
from src.database.review_manager import review_record


DECISIONS = {"approved", "rejected", "corrected"}

CORRECTION_COLUMNS = {
    "Corrected Year": "year",
    "Corrected Category": "category",
    "Corrected Metric Type": "metric_type",
    "Corrected Count": "count",
}

ORIGINAL_COLUMNS = {
    "Year": "year",
    "Category": "category",
    "Metric Type": "metric_type",
    "Count": "count",
}


def is_empty(value):
    return value is None or (
        isinstance(value, str) and not value.strip()
    )


def import_review_sheet(
    excel_path: Path,
    connection
) -> dict:

    workbook = load_workbook(excel_path, data_only=True)

    try:
        sheet = workbook["Review Records"]

        headers = {
            cell.value: index
            for index, cell in enumerate(sheet[1])
            if cell.value is not None
        }

        required = {
            "Record ID",
            "Filename",
            "Review Decision",
            "Reviewer",
            "Reviewer Notes",
            *ORIGINAL_COLUMNS,
            *CORRECTION_COLUMNS,
        }

        missing = required - set(headers)

        if missing:
            raise ValueError(
                f"Missing Excel columns: {sorted(missing)}"
            )

        summary = {
            "approved": 0,
            "rejected": 0,
            "corrected": 0,
            "skipped": 0,
        }

        seen_ids = set()

        def cell_value(row, name):
            return row[headers[name]]

        # One transaction for the entire workbook.
        with connection:
            for excel_row_number, row in enumerate(
                sheet.iter_rows(min_row=2, values_only=True),
                start=2
            ):

                decision = cell_value(row, "Review Decision")

                # Unreviewed rows remain pending.
                if is_empty(decision):
                    summary["skipped"] += 1
                    continue

                decision = str(decision).strip().lower()

                if decision not in DECISIONS:
                    raise ValueError(
                        f"Row {excel_row_number}: "
                        f"Invalid decision '{decision}'"
                    )

                record_id = cell_value(row, "Record ID")

                if type(record_id) is not int or record_id <= 0:
                    raise ValueError(
                        f"Row {excel_row_number}: Invalid record ID"
                    )

                if record_id in seen_ids:
                    raise ValueError(
                        f"Row {excel_row_number}: Duplicate record ID"
                    )

                seen_ids.add(record_id)

                reviewer = cell_value(row, "Reviewer")

                if is_empty(reviewer):
                    raise ValueError(
                        f"Row {excel_row_number}: Reviewer required"
                    )

                reviewer = str(reviewer).strip()

                notes = cell_value(row, "Reviewer Notes")
                notes = (
                    None if is_empty(notes)
                    else str(notes).strip()
                )

                record = connection.execute(
                    """
                    SELECT
                        er.id,
                        er.year,
                        er.category,
                        er.metric_type,
                        er.count,
                        er.review_status,
                        sf.filename
                    FROM extraction_records er
                    JOIN extraction_runs run
                        ON er.extraction_run_id = run.id
                    JOIN source_files sf
                        ON run.source_file_id = sf.id
                    WHERE er.id = ?
                      AND run.status = 'completed'
                      AND run.model_provider != 'gemini_test'
                    """,
                    (record_id,)
                ).fetchone()

                if record is None:
                    raise ValueError(
                        f"Row {excel_row_number}: "
                        f"Record {record_id} not found"
                    )

                # Verify the Excel row belongs to this record.
                if record["filename"] != cell_value(row, "Filename"):
                    raise ValueError(
                        f"Row {excel_row_number}: Filename mismatch"
                    )

                for excel_column, db_field in ORIGINAL_COLUMNS.items():
                    excel_value = cell_value(row, excel_column)
                    database_value = record[db_field]

                    if excel_value != database_value:
                        raise ValueError(
                            f"Row {excel_row_number}: "
                            f"Original {excel_column} was changed"
                        )

                corrections = {}

                for excel_column, field in CORRECTION_COLUMNS.items():
                    value = cell_value(row, excel_column)

                    if not is_empty(value):
                        if field == "count":
                            if (
                                type(value) is not int
                                or value < 0
                            ):
                                raise ValueError(
                                    f"Row {excel_row_number}: "
                                    "Corrected count must be "
                                    "a nonnegative integer"
                                )
                        else:
                            if not isinstance(value, str):
                                raise ValueError(
                                    f"Row {excel_row_number}: "
                                    f"{excel_column} must be text"
                                )

                            value = value.strip()

                        corrections[field] = value

                if decision == "corrected":
                    if not corrections:
                        raise ValueError(
                            f"Row {excel_row_number}: "
                            "Correction values required"
                        )

                    # Validate the complete corrected record
                    # against our existing Pydantic schema.
                    candidate = {
                        "year": record["year"],
                        "category": record["category"],
                        "metric_type": record["metric_type"],
                        "count": record["count"],
                        "evidence_text": "Review validation"
                    }

                    candidate.update(corrections)

                    ExtractionRecord.model_validate(candidate)

                elif corrections:
                    raise ValueError(
                        f"Row {excel_row_number}: "
                        "Remove correction values unless "
                        "decision is corrected"
                    )

                if record["review_status"] != "pending":
                    # Already reviewed: don't import it again.
                    # Different decisions also require an
                    # explicit review-revision workflow.
                    raise ValueError(
                        f"Row {excel_row_number}: "
                        f"Record {record_id} was already reviewed"
                    )

                review_record(
                    connection=connection,
                    record_id=record_id,
                    decision=decision,
                    reviewer=reviewer,
                    notes=notes,
                    corrections=corrections
                )

                summary[decision] += 1

        return summary

    finally:
        workbook.close()
