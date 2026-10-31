import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = [
    "Record ID", "Council", "Filename", "Year",
    "Category", "Metric Type", "Count", "Evidence",
    "Source Location", "Review Decision",
    "Corrected Year", "Corrected Category",
    "Corrected Metric Type", "Corrected Count",
    "Reviewer", "Reviewer Notes"
]


def get_extracted_rows(extracted_directory):
    rows = []
    seen_hashes = set()

    for path in sorted(Path(extracted_directory).glob("*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        file_hash = data["source"]["file_hash"]
        if file_hash in seen_hashes:
            continue

        seen_hashes.add(file_hash)
        extraction = data["extraction"]

        for index, record in enumerate(extraction["records"], start=1):
            rows.append([
                f"{file_hash}-{index:04d}",
                extraction.get("council_name"),
                data.get("filename"),
                record.get("year"),
                record["category"],
                record["metric_type"],
                record["count"],
                record["evidence_text"],
                record.get("source_location"),
                None, None, None, None, None, None, None
            ])

    return rows

def format_review_sheet(sheet, start_row=2):
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions

    widths = [
        75, 30, 30, 15, 30, 26, 12, 65,
        35, 20, 18, 30, 26, 20, 22, 45
    ]

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width

    yellow = PatternFill(fill_type="solid", fgColor="FFF2CC")
    for row in sheet.iter_rows(min_row=start_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        for cell in row[9:16]:
            cell.fill = yellow

    if sheet.max_row >= 2: 
        validation = DataValidation(type="list", formula1='"approved,rejected,corrected"', allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(f"J2:J{sheet.max_row}")


def export_review_sheet(extracted_directory, output_path):
    output_path = Path(output_path)

    if output_path.exists():
        raise FileExistsError(output_path)

    rows = get_extracted_rows(extracted_directory)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Review Records"
    sheet.append(HEADERS)

    for row in rows:
        sheet.append(row)

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="244062")

    format_review_sheet(sheet)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    workbook.close()

    return len(rows)
