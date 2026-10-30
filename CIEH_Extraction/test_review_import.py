
import sqlite3
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from src.database.connection import get_connection
from src.review.import_review_sheet import import_review_sheet


BASE_DIR = Path(__file__).resolve().parent
EXCEL_PATH = BASE_DIR / "data" / "review" / "cieh_review.xlsx"


def test_import():
    if not EXCEL_PATH.exists():
        raise FileNotFoundError(EXCEL_PATH)

    with tempfile.TemporaryDirectory() as temp_dir:
        temp_dir = Path(temp_dir)

        test_db = temp_dir / "test.db"
        test_excel = temp_dir / "test_review.xlsx"

        # Copy SQLite safely using its backup API.
        source_connection = get_connection()
        destination_connection = sqlite3.connect(test_db)

        try:
            source_connection.backup(destination_connection)
        finally:
            source_connection.close()
            destination_connection.close()

        # Open the original Excel and save a temporary copy.
        workbook = load_workbook(EXCEL_PATH)
        sheet = workbook["Review Records"]

        headers = {
            cell.value: index
            for index, cell in enumerate(sheet[1], start=1)
        }

        def set_value(row, column, value):
            sheet.cell(
                row=row,
                column=headers[column]
            ).value = value

        # Use first three genuine records.
        if sheet.max_row < 4:
            raise ValueError("Need at least three records")

        set_value(2, "Review Decision", "approved")
        set_value(2, "Reviewer", "Test Reviewer")

        set_value(3, "Review Decision", "corrected")
        set_value(3, "Reviewer", "Test Reviewer")
        set_value(
            3,
            "Corrected Count",
            sheet.cell(
                row=3,
                column=headers["Count"]
            ).value + 1
        )

        set_value(4, "Review Decision", "rejected")
        set_value(4, "Reviewer", "Test Reviewer")
        set_value(4, "Reviewer Notes", "Test rejection")

        workbook.save(test_excel)
        workbook.close()

        connection = sqlite3.connect(test_db)
        connection.row_factory = sqlite3.Row
        connection.execute("PRAGMA foreign_keys = ON")

        try:
            summary = import_review_sheet(
                test_excel,
                connection
            )

            print("First import:", summary)

            assert summary["approved"] == 1
            assert summary["corrected"] == 1
            assert summary["rejected"] == 1

            results = connection.execute("""
                SELECT decision, COUNT(*) AS total
                FROM record_reviews
                GROUP BY decision
            """).fetchall()

            print(
                "Saved reviews:",
                {r["decision"]: r["total"] for r in results}
            )

            # A repeated import must not create duplicates.
            try:
                import_review_sheet(test_excel, connection)
            except ValueError as error:
                print("Repeated import blocked:", error)
            else:
                raise AssertionError(
                    "Repeated import was not blocked"
                )

            total_reviews = connection.execute(
                "SELECT COUNT(*) FROM record_reviews"
            ).fetchone()[0]

            assert total_reviews == 3

            print("All tests passed!")

        finally:
            connection.close()


if __name__ == "__main__":
    test_import()
