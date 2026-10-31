import os
import shutil
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook
from src.review.export_review_sheet import (
    export_review_sheet,
    get_extracted_rows,
    format_review_sheet
)

def sync_review_sheet(extracted_directory, review_path):
    review_path = Path(review_path)
    if not review_path.exists():
        total = export_review_sheet(extracted_directory,review_path)
        return {
            "added": total,
            "existing": 0,
            "total": total
        }

    workbook = load_workbook(review_path)

    try:
        sheet = workbook["Review Records"]
        existing_ids = set()
        for row in sheet.iter_rows(min_row=2, values_only=True):
            record_id = row[0]
            if record_id:
                if record_id in existing_ids:
                    raise ValueError(f"Duplicate record ID: {record_id}")

                existing_ids.add(record_id)

        existing_count = len(existing_ids)
        new_rows = get_extracted_rows(extracted_directory)
        added = 0

        for row in new_rows:
            if row[0] in existing_ids:
                continue

            sheet.append(row)
            existing_ids.add(row[0])
            added += 1

        if added:
            format_review_sheet(sheet)
            backup_dir = review_path.parent / "backups"
            backup_dir.mkdir(parents=True, exist_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            backup_path = (backup_dir / f"review_{timestamp}.xlsx")
            shutil.copy2(review_path, backup_path)
            temporary_path = review_path.with_name(review_path.stem + "_updating.xlsx")

            try:
                workbook.save(temporary_path)
                os.replace(temporary_path, review_path)
            finally:
                temporary_path.unlink(missing_ok=True)

        return {
            "added": added,
            "existing": existing_count,
            "total": len(existing_ids)
        }

    finally:
        workbook.close()
