import csv
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from src.ai.extraction_schema import ExtractionRecord


HEADERS = [
    "Record ID", "Council", "Filename", "Year",
    "Category", "Metric Type", "Count", "Evidence",
    "Source Location", "Review Decision",
    "Reviewer", "Reviewer Notes"
]

CORRECTIONS = {
    "Year": "Corrected Year",
    "Category": "Corrected Category",
    "Metric Type": "Corrected Metric Type",
    "Count": "Corrected Count"
}


def safe_excel_value(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value

    return value


def read_verified_records(review_path):
    workbook = load_workbook(review_path, read_only=True, data_only=True)

    try:
        sheet = workbook["Review Records"]
        rows = sheet.iter_rows(values_only=True)
        columns = {
            name: index
            for index, name in enumerate(next(rows))
        }
        required = set(HEADERS) | set(CORRECTIONS.values())
        missing = required - columns.keys()
        if missing:
            raise ValueError(f"Missing columns: {sorted(missing)}")

        verified = []
        seen_ids = set()
        for row_number, row in enumerate(rows, start=2):
            def get(column):
                return row[columns[column]]
            record_id = get("Record ID")

            if not record_id:
                continue

            if record_id in seen_ids:
                raise ValueError(f"Row {row_number}: Duplicate record ID")
            seen_ids.add(record_id)
            decision = str(get("Review Decision") or "").strip().lower()

            if not decision:
                continue

            if decision not in {"approved", "rejected", "corrected"}:
                raise ValueError(f"Row {row_number}: Invalid review decision")

            if decision == "rejected":
                continue

            reviewer = str(get("Reviewer") or "").strip()
            if not reviewer:
                raise ValueError(f"Row {row_number}: Reviewer is required")

            corrections = {
                field: get(column)
                for field, column in CORRECTIONS.items()
                if get(column) is not None
                and str(get(column)).strip() != ""
            }

            if decision == "approved" and corrections:
                raise ValueError(f"Row {row_number}: Approved record "
                    "cannot contain corrections"
                )

            if decision == "corrected" and not corrections:
                raise ValueError(f"Row {row_number}: Correction required")

            final = {
                field: corrections.get(field, get(field))
                for field in CORRECTIONS
            }

            validated = ExtractionRecord.model_validate({
                "year": final["Year"],
                "category": final["Category"],
                "metric_type": final["Metric Type"],
                "count": final["Count"],
                "evidence_text": get("Evidence"),
                "source_location": get("Source Location")
            })

            record = {
                column: get(column)
                for column in HEADERS
            }

            record.update({
                "Year": validated.year,
                "Category": validated.category,
                "Metric Type": validated.metric_type,
                "Count": validated.count,
                "Review Decision": decision,
                "Reviewer": reviewer
            })
            verified.append(record)

        return verified

    finally:
        workbook.close()


def export_verified_data(review_path, output_directory):
    records = read_verified_records(review_path)
    output_directory = Path(output_directory)
    output_directory.mkdir(parents=True, exist_ok=True)
    csv_path = output_directory / "cieh_verified.csv"
    excel_path = output_directory / "cieh_verified.xlsx"

    with csv_path.open("w", newline="", encoding="utf-8-sig") as file:
        writer = csv.writer(file)
        writer.writerow(HEADERS)
        for record in records:
            writer.writerow([safe_excel_value(record[column]) for column in HEADERS])

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Verified Data"
    sheet.append(HEADERS)

    for record in records:
        sheet.append([safe_excel_value(record[column])for column in HEADERS])

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid",fgColor="244062")

    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions

    widths = [
        35, 28, 30, 15, 30, 26,
        12, 65, 35, 20, 22, 45
    ]

    for column, width in zip("ABCDEFGHIJKL", widths):
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top",wrap_text=True)

    workbook.save(excel_path)
    workbook.close()
    print(f"Verified records: {len(records)}")
    print(f"CSV: {csv_path}")
    print(f"Excel: {excel_path}")
    return len(records)
